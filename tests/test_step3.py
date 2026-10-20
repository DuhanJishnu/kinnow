from pathlib import Path
import docx
import openpyxl
from PIL import Image
import pptx
import pymupdf
import pytest

from dropwheel.converters import load_all, registry


@pytest.fixture(autouse=True)
def init_all():
    load_all()


def _make_sample_pptx(path: Path) -> None:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Kinnow Presentation"
    slide.placeholders[1].text = "Automated slide conversion test"
    prs.save(str(path))


def _make_sample_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=600, height=400)
    page.insert_text((50, 50), "PDF Title Header\nItem\tQty\nKinnow\t42")
    doc.save(str(path))
    doc.close()


def test_pptx_to_pdf(tmp_path: Path):
    src = tmp_path / "slides.pptx"
    _make_sample_pptx(src)

    out = registry.convert(src, "pdf")
    assert out.exists() and out.suffix == ".pdf"
    with pymupdf.open(out) as doc:
        assert doc.page_count >= 1
        text = doc[0].get_text()
        assert "Kinnow Presentation" in text


def test_pptx_to_png(tmp_path: Path):
    src = tmp_path / "slides.pptx"
    _make_sample_pptx(src)

    out = registry.convert(src, "png")
    assert out.exists()
    target_file = out if out.is_file() else (out / "slide_1.png")
    assert target_file.exists()
    with Image.open(target_file) as im:
        assert im.width > 50


def test_pptx_to_jpg(tmp_path: Path):
    src = tmp_path / "slides.pptx"
    _make_sample_pptx(src)

    out = registry.convert(src, "jpg")
    assert out.exists()
    target_file = out if out.is_file() else (out / "slide_1.jpg")
    assert target_file.exists()


def test_pdf_to_md(tmp_path: Path):
    src = tmp_path / "doc.pdf"
    _make_sample_pdf(src)

    out = registry.convert(src, "md")
    assert out.exists() and out.suffix == ".md"
    content = out.read_text(encoding="utf-8")
    assert "PDF Title Header" in content


def test_pdf_to_docx(tmp_path: Path):
    src = tmp_path / "doc.pdf"
    _make_sample_pdf(src)

    out = registry.convert(src, "docx")
    assert out.exists() and out.suffix == ".docx"
    doc = docx.Document(str(out))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "PDF Title Header" in full_text or "Kinnow" in full_text


def test_pdf_to_xlsx(tmp_path: Path):
    src = tmp_path / "doc.pdf"
    _make_sample_pdf(src)

    out = registry.convert(src, "xlsx")
    assert out.exists() and out.suffix == ".xlsx"
    wb = openpyxl.load_workbook(str(out))
    ws = wb.active
    values = [str(cell) for row in ws.iter_rows(values_only=True) for cell in row if cell is not None]
    assert any("Kinnow" in v for v in values)


def test_pdf_to_pptx(tmp_path: Path):
    src = tmp_path / "doc.pdf"
    _make_sample_pdf(src)

    out = registry.convert(src, "pptx")
    assert out.exists() and out.suffix == ".pptx"
    prs = pptx.Presentation(str(out))
    assert len(prs.slides) == 1


def test_pdf_to_jpg(tmp_path: Path):
    src = tmp_path / "doc.pdf"
    _make_sample_pdf(src)

    out = registry.convert(src, "jpg")
    assert out.exists()
    target_file = out if out.is_file() else (out / "page_1.jpg")
    assert target_file.exists()
