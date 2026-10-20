from pathlib import Path
import zipfile
import docx
import openpyxl
from PIL import Image
import pymupdf
import pytest

from dropwheel.converters import load_all, registry


@pytest.fixture(autouse=True)
def init_all():
    load_all()


def _make_sample_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_heading("Project Kinnow", 0)
    doc.add_paragraph("First paragraph testing docx conversion.")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    doc.save(str(path))


def _make_sample_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Item", "Price", "Qty"])
    ws.append(["Apple", 1.5, 10])
    ws.append(["Kinnow", 2.0, 25])
    wb.save(str(path))


def test_docx_to_txt(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "txt")
    assert out.exists() and out.suffix == ".txt"
    content = out.read_text(encoding="utf-8")
    assert "Project Kinnow" in content
    assert "First paragraph testing docx" in content


def test_docx_to_html(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "html")
    assert out.exists() and out.suffix == ".html"
    content = out.read_text(encoding="utf-8")
    assert "Project Kinnow" in content
    assert "<p>" in content or "<h1>" in content


def test_docx_to_md(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "md")
    assert out.exists() and out.suffix == ".md"
    content = out.read_text(encoding="utf-8")
    assert "Project Kinnow" in content


def test_docx_to_pdf(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "pdf")
    assert out.exists() and out.suffix == ".pdf"
    with pymupdf.open(out) as doc:
        assert doc.page_count >= 1
        assert "Project Kinnow" in doc[0].get_text()


def test_docx_to_png(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "png")
    assert out.exists() and out.suffix == ".png"
    with Image.open(out) as im:
        assert im.width > 100
        assert im.height > 100


def test_docx_to_epub(tmp_path: Path):
    src = tmp_path / "sample.docx"
    _make_sample_docx(src)

    out = registry.convert(src, "epub")
    assert out.exists() and out.suffix == ".epub"
    with zipfile.ZipFile(out, "r") as zf:
        assert "mimetype" in zf.namelist()
        assert zf.read("mimetype") == b"application/epub+zip"
        assert "OEBPS/content.opf" in zf.namelist()


def test_sheet_to_csv(tmp_path: Path):
    src = tmp_path / "data.xlsx"
    _make_sample_xlsx(src)

    out = registry.convert(src, "csv")
    assert out.exists() and out.suffix == ".csv"
    content = out.read_text(encoding="utf-8-sig")
    assert "Item,Price,Qty" in content
    assert "Kinnow" in content
    assert "25" in content


def test_sheet_to_tsv(tmp_path: Path):
    src = tmp_path / "data.xlsx"
    _make_sample_xlsx(src)

    out = registry.convert(src, "tsv")
    assert out.exists() and out.suffix == ".tsv"
    content = out.read_text(encoding="utf-8-sig")
    assert "Item\tPrice\tQty" in content
    assert "Kinnow" in content
    assert "25" in content


def test_sheet_to_pdf(tmp_path: Path):
    src = tmp_path / "data.xlsx"
    _make_sample_xlsx(src)

    out = registry.convert(src, "pdf")
    assert out.exists() and out.suffix == ".pdf"
    with pymupdf.open(out) as doc:
        assert doc.page_count >= 1
        text = doc[0].get_text()
        assert "Kinnow" in text


def test_sheet_to_ods(tmp_path: Path):
    src = tmp_path / "data.xlsx"
    _make_sample_xlsx(src)

    out = registry.convert(src, "ods")
    assert out.exists() and out.suffix == ".ods"
    with zipfile.ZipFile(out, "r") as zf:
        assert "mimetype" in zf.namelist()
        assert zf.read("mimetype") == b"application/vnd.oasis.opendocument.spreadsheet"
