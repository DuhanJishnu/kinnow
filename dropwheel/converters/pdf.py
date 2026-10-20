"""PDF conversions using PyMuPDF (imported as `fitz` / `pymupdf`), pdf2docx, and openpyxl."""
from __future__ import annotations

import io
from pathlib import Path

from .registry import register
from dropwheel.utils import unique_path


@register(["pdf"], "png")
def pdf_to_png(src: Path, dst: Path) -> Path:
    import pymupdf

    with pymupdf.open(src) as doc:
        if doc.page_count <= 1:
            if doc.page_count == 1:
                doc[0].get_pixmap(dpi=150).save(str(dst))
            return dst

        folder = unique_path(dst.with_suffix(""))
        folder.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(doc, start=1):
            page.get_pixmap(dpi=150).save(str(folder / f"page_{i}.png"))
        return folder


@register(["pdf"], "jpg")
def pdf_to_jpg(src: Path, dst: Path) -> Path:
    import pymupdf

    with pymupdf.open(src) as doc:
        if doc.page_count <= 1:
            if doc.page_count == 1:
                doc[0].get_pixmap(dpi=150).save(str(dst))
            return dst

        folder = unique_path(dst.with_suffix(""))
        folder.mkdir(parents=True, exist_ok=True)
        for i, page in enumerate(doc, start=1):
            page.get_pixmap(dpi=150).save(str(folder / f"page_{i}.jpg"))
        return folder


@register(["pdf"], "txt")
def pdf_to_txt(src: Path, dst: Path) -> None:
    import pymupdf

    with pymupdf.open(src) as doc:
        text = "\n\n".join(page.get_text() for page in doc)
    dst.write_text(text, encoding="utf-8")


@register(["pdf"], "md")
def pdf_to_md(src: Path, dst: Path) -> None:
    """Extract structured Markdown from PDF."""
    import pymupdf

    lines = []
    with pymupdf.open(src) as doc:
        for i, page in enumerate(doc, start=1):
            if doc.page_count > 1:
                lines.append(f"\n## Page {i}\n")
            text = page.get_text("text")
            lines.append(text.strip())

    dst.write_text("\n\n".join(l for l in lines if l), encoding="utf-8")


@register(["pdf"], "docx")
def pdf_to_docx(src: Path, dst: Path) -> None:
    """Convert PDF layout and content to editable Microsoft Word (.docx)."""
    from pdf2docx import Converter

    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    finally:
        cv.close()


@register(["pdf"], "xlsx")
def pdf_to_xlsx(src: Path, dst: Path) -> None:
    """Extract tables and text rows from PDF into an Excel spreadsheet (.xlsx)."""
    import openpyxl
    import pymupdf

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Page 1"

    with pymupdf.open(src) as doc:
        for page_idx, page in enumerate(doc):
            sheet = ws if page_idx == 0 else wb.create_sheet(title=f"Page {page_idx + 1}")

            tables = page.find_tables()
            has_table = False
            for t in tables:
                for row in t.extract():
                    if any(cell is not None for cell in row):
                        sheet.append([str(c) if c is not None else "" for c in row])
                        has_table = True
                sheet.append([])

            if not has_table:
                text = page.get_text("text")
                for line in text.splitlines():
                    if line.strip():
                        parts = line.split("\t") if "\t" in line else [line.strip()]
                        sheet.append(parts)

    wb.save(str(dst))


@register(["pdf"], "pptx")
def pdf_to_pptx(src: Path, dst: Path) -> None:
    """Convert PDF pages to PowerPoint presentation slides (.pptx)."""
    import pptx
    from pptx.util import Inches
    import pymupdf

    prs = pptx.Presentation()
    blank_layout = prs.slide_layouts[6]

    with pymupdf.open(src) as doc:
        for page in doc:
            prs.slide_width = Inches(page.rect.width / 72.0)
            prs.slide_height = Inches(page.rect.height / 72.0)

            slide = prs.slides.add_slide(blank_layout)
            pix = page.get_pixmap(dpi=150)
            img_bytes = io.BytesIO(pix.tobytes("png"))
            slide.shapes.add_picture(
                img_bytes, 0, 0, width=prs.slide_width, height=prs.slide_height
            )

    prs.save(str(dst))
